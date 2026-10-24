"""Fase 5: itens de orçamento vindos da base própria do usuário (fonte = PROPRIA)."""

import io
from unittest.mock import MagicMock

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.dependencies import get_supabase
from app.main import app
from app.modules.importacao.services.pdf_service import PdfService
from app.modules.orcamento.schemas import OrcamentoItemCreate
from app.modules.orcamento.services import OrcamentoItemService, anexar_nome_base_propria
from core.security import get_current_user
from tests.test_base_propria import FakeSupabase, item, nova_composicao, novo_insumo, semear_precos_oficiais

USUARIO_A = "user-a"
USUARIO_B = "user-b"
MES = "10/2026"
TIPO = "Sem Desoneração"


@pytest.fixture
def banco():
    return FakeSupabase()


@pytest.fixture
def api(banco):
    app.dependency_overrides[get_supabase] = lambda: banco

    def como(usuario):
        app.dependency_overrides[get_current_user] = lambda: {"id": usuario, "email": f"{usuario}@x.com"}
        return TestClient(app)

    yield como
    app.dependency_overrides.clear()


@pytest.fixture
def a(api):
    return api(USUARIO_A)


def criar_orcamento(banco, dono=USUARIO_A, bdi=10.0):
    banco.tabelas.setdefault("orcamentos", []).append(
        {
            "id": f"orc-{dono}",
            "user_id": dono,
            "nome": "Obra X",
            "cliente": "Cliente",
            "estado": "CE",
            "base_referencia": MES,
            "tipo_composicao": TIPO,
            "fonte": "SINAPI",
            "bdi": bdi,
            "bdi_config": None,
            "created_at": "2026-01-01T00:00:00",
        }
    )
    return f"orc-{dono}"


@pytest.fixture
def cenario(a, banco):
    """Base do usuário A com C1 = 2 x INS-1 (10) + 1 x SINAPI 370 (50) = 70, e um orçamento do mesmo usuário."""
    semear_precos_oficiais(banco)
    base = a.post("/bases-proprias", json={"nome": "Minha base", "uf": "ce"}).json()["id"]
    insumo = novo_insumo(a, base, "INS-1", 10, tipo="MAO_DE_OBRA")
    nova_composicao(a, base, "C1", [item("PROPRIA", "INSUMO", "INS-1", 2), item("SINAPI", "INSUMO", "370", 1)])
    return {"base": base, "insumo": insumo, "orcamento": criar_orcamento(banco)}


def payload(base, codigo="C1", quantidade=3, **extra):
    return {
        "codigo_composicao": codigo,
        "descricao": "Serviço próprio",
        "quantidade": quantidade,
        "unidade": "M2",
        "fonte": "PROPRIA",
        "base_propria_id": base,
        **extra,
    }


def incluir(a, c, **extra):
    r = a.post(f"/orcamentos/{c['orcamento']}/itens", json=payload(c["base"], **extra))
    assert r.status_code == 200, r.text
    return r.json()


# --- Inclusão ----------------------------------------------------------------


def test_inclui_item_proprio_com_preco_de_insumo_sinapi_dentro_da_composicao(a, banco, cenario):
    novo = incluir(a, cenario)

    assert novo["fonte"] == "PROPRIA"
    assert novo["base_propria_id"] == cenario["base"]
    assert novo["preco_unitario"] == 70.0  # 2 x 10 (própria) + 1 x 50 (SINAPI 370 no mês do orçamento)
    assert novo["preco_total"] == 210.0
    assert novo["preco_unitario_bdi"] == 77.0
    assert novo["preco_total_bdi"] == 231.0
    assert banco.tabelas["orcamentos"][0]["valor_total"] == 231.0


def test_preenche_analitico_para_curva_abc(a, banco, cenario):
    novo = incluir(a, cenario)

    insumos = {i["codigo_insumo"]: i for i in banco.tabelas["orcamento_item_insumo"]}
    assert set(insumos) == {"INS-1", "370"}
    assert insumos["INS-1"]["orcamento_item_id"] == novo["id"]
    assert insumos["INS-1"]["quantidade_unitaria"] == 6.0
    assert insumos["INS-1"]["total"] == 60.0
    assert insumos["INS-1"]["tipo_item"] == "MAO_DE_OBRA"
    assert insumos["370"]["total"] == 150.0
    assert insumos["370"]["tipo_item"] == "MATERIAL"

    curva = a.get(f"/orcamentos/{cenario['orcamento']}/curva-abc")
    assert curva.status_code == 200
    corpo = curva.json()
    assert corpo["valor_total"] == 210.0
    assert [i["codigo_insumo"] for i in corpo["insumos"]] == ["370", "INS-1"]


def test_preco_fica_congelado_quando_a_base_muda(a, banco, cenario):
    incluir(a, cenario)
    r = a.put(f"/bases-proprias/{cenario['base']}/insumos/{cenario['insumo']['id']}", json={"preco": 99})
    assert r.status_code == 200

    itens = a.get(f"/orcamentos/{cenario['orcamento']}/itens").json()
    assert itens[0]["preco_unitario"] == 70.0
    assert itens[0]["base_propria_nome"] == "Minha base"


def test_preco_informado_pelo_cliente_prevalece(a, cenario):
    novo = incluir(a, cenario, preco_unitario=65.5)
    assert novo["preco_unitario"] == 65.5


def test_item_proprio_exige_base_e_codigo_existentes(a, banco, cenario):
    url = f"/orcamentos/{cenario['orcamento']}/itens"
    sem_base = {**payload(cenario["base"]), "base_propria_id": None}
    assert a.post(url, json=sem_base).status_code == 400
    assert a.post(url, json=payload(cenario["base"], codigo="NAO-EXISTE")).status_code == 404
    assert a.post(url, json={**payload(cenario["base"]), "codigo_composicao": "MANUAL"}).status_code == 400
    assert banco.tabelas.get("orcamento_itens", []) == []


def test_composicao_com_item_sem_preco_nao_entra_sem_preco_informado(a, banco, cenario):
    nova_composicao(a, cenario["base"], "C2", [item("SINAPI", "INSUMO", "NAO-EXISTE", 1)])
    r = a.post(f"/orcamentos/{cenario['orcamento']}/itens", json=payload(cenario["base"], codigo="C2"))
    assert r.status_code == 400
    assert "NAO-EXISTE" in r.json()["detail"]

    assert incluir(a, cenario, codigo="C2", preco_unitario=12)["preco_unitario"] == 12


def test_coluna_nova_so_vai_no_insert_de_item_proprio(a, banco, cenario):
    banco.tabelas["orcamento_itens"] = []
    incluir(a, cenario)
    assert "base_propria_id" in banco.tabelas["orcamento_itens"][0]


# --- Isolamento --------------------------------------------------------------


def test_nao_usa_base_de_outro_usuario(api, banco, cenario):
    b = api(USUARIO_B)
    orc_b = criar_orcamento(banco, dono=USUARIO_B)

    r = b.post(f"/orcamentos/{orc_b}/itens", json=payload(cenario["base"]))
    assert r.status_code == 404
    assert "Base própria não encontrada" in r.json()["detail"]
    assert banco.tabelas.get("orcamento_itens", []) == []


def test_nao_acessa_orcamento_de_outro_usuario(api, a, banco, cenario):
    incluir(a, cenario)
    b = api(USUARIO_B)
    base_url = f"/orcamentos/{cenario['orcamento']}"
    assert b.post(f"{base_url}/itens/atualizar-precos-propria").status_code == 404
    assert b.get(f"{base_url}/buscar-composicoes-proprias?termo=C1").status_code == 404
    assert b.get(f"{base_url}/itens").status_code == 404


def test_busca_so_mostra_bases_do_dono_do_orcamento(api, banco, cenario):
    b = api(USUARIO_B)
    orc_b = criar_orcamento(banco, dono=USUARIO_B)
    assert b.get(f"/orcamentos/{orc_b}/buscar-composicoes-proprias?termo=C1").json() == []


# --- Atualização de preços ---------------------------------------------------


def encarecer_insumo(a, c, preco):
    r = a.put(f"/bases-proprias/{c['base']}/insumos/{c['insumo']['id']}", json={"preco": preco})
    assert r.status_code == 200


def test_previa_nao_grava_e_aplicar_atualiza_preco_e_totais(a, banco, cenario):
    incluir(a, cenario)
    encarecer_insumo(a, cenario, 15)  # C1 passa de 70 para 80
    url = f"/orcamentos/{cenario['orcamento']}/itens/atualizar-precos-propria"

    previa = a.post(url)
    assert previa.status_code == 200
    corpo = previa.json()
    assert corpo["aplicado"] is False
    assert corpo["itens_alterados"] == 1
    assert corpo["diferenca_custo_direto"] == 30.0  # (80 - 70) x 3
    linha = corpo["itens"][0]
    assert (linha["preco_atual"], linha["preco_novo"], linha["diferenca"]) == (70.0, 80.0, 10.0)
    assert linha["variacao_percentual"] == 14.29
    assert linha["aplicado"] is False
    assert banco.tabelas["orcamento_itens"][0]["preco_unitario"] == 70.0
    assert banco.tabelas["orcamentos"][0]["valor_total"] == 231.0

    aplicado = a.post(f"{url}?aplicar=true").json()
    assert aplicado["aplicado"] is True
    assert aplicado["itens"][0]["aplicado"] is True
    gravado = banco.tabelas["orcamento_itens"][0]
    assert gravado["preco_unitario"] == 80.0
    assert gravado["preco_total"] == 240.0
    assert gravado["preco_unitario_bdi"] == 88.0
    assert gravado["preco_total_bdi"] == 264.0
    assert banco.tabelas["orcamentos"][0]["valor_total"] == 264.0
    insumos = {i["codigo_insumo"]: i for i in banco.tabelas["orcamento_item_insumo"]}
    assert insumos["INS-1"]["total"] == 90.0
    assert len(banco.tabelas["orcamento_item_insumo"]) == 2

    de_novo = a.post(url).json()
    assert de_novo["itens_alterados"] == 0
    assert de_novo["itens"][0]["alterado"] is False


def test_atualizacao_ignora_item_sem_base_composicao_removida_e_insumo_editado(a, banco, cenario):
    primeiro = incluir(a, cenario)
    segundo = incluir(a, cenario)
    terceiro = incluir(a, cenario)
    orfao = incluir(a, cenario)
    for linha in banco.tabelas["orcamento_itens"]:
        if linha["id"] == orfao["id"]:
            linha["base_propria_id"] = None  # base excluída: ON DELETE SET NULL
        if linha["id"] == segundo["id"]:
            linha["codigo_composicao"] = "FANTASMA"
    for ins in banco.tabelas["orcamento_item_insumo"]:
        if ins["orcamento_item_id"] == terceiro["id"] and ins["codigo_insumo"] == "INS-1":
            ins["preco_unitario_custom"] = 11.0
    encarecer_insumo(a, cenario, 15)

    corpo = a.post(f"/orcamentos/{cenario['orcamento']}/itens/atualizar-precos-propria?aplicar=true").json()

    assert [i["item_id"] for i in corpo["itens"]] == [primeiro["id"]]
    motivos = {i["item_id"]: i["motivo"] for i in corpo["ignorados"]}
    assert "não existe mais" in motivos[segundo["id"]]
    assert "editados" in motivos[terceiro["id"]]
    assert "sem vínculo" in motivos[orfao["id"]]
    por_id = {i["id"]: i for i in banco.tabelas["orcamento_itens"]}
    assert por_id[primeiro["id"]]["preco_unitario"] == 80.0
    assert por_id[terceiro["id"]]["preco_unitario"] == 70.0


def test_atualizacao_ignora_item_cuja_base_virou_de_outro_dono(a, banco, cenario):
    incluir(a, cenario)
    banco.tabelas["bases_proprias"][0]["user_id"] = USUARIO_B
    corpo = a.post(f"/orcamentos/{cenario['orcamento']}/itens/atualizar-precos-propria?aplicar=true").json()
    assert corpo["itens"] == []
    assert corpo["ignorados"][0]["motivo"] == "Base própria não encontrada"


def test_atualizacao_sem_itens_proprios(a, banco, cenario):
    corpo = a.post(f"/orcamentos/{cenario['orcamento']}/itens/atualizar-precos-propria?aplicar=true").json()
    assert corpo == {
        "aplicado": True,
        "total_itens": 0,
        "itens_alterados": 0,
        "diferenca_custo_direto": 0.0,
        "itens": [],
        "ignorados": [],
    }


def test_atualizacao_nao_toca_itens_sinapi(a, banco, cenario):
    incluir(a, cenario)
    banco.tabelas["orcamento_itens"].append(
        {"id": "sinapi-1", "orcamento_id": cenario["orcamento"], "fonte": "SINAPI", "preco_unitario": 5.0}
    )
    encarecer_insumo(a, cenario, 15)
    corpo = a.post(f"/orcamentos/{cenario['orcamento']}/itens/atualizar-precos-propria?aplicar=true").json()
    assert corpo["total_itens"] == 1
    assert banco.tabelas["orcamento_itens"][1]["preco_unitario"] == 5.0


# --- Edição de item ----------------------------------------------------------


def test_edita_quantidade_de_item_proprio_e_recusa_trocar_composicao(a, banco, cenario):
    novo = incluir(a, cenario)
    url = f"/orcamentos/{cenario['orcamento']}/itens/{novo['id']}"

    ok = a.put(url, json={"quantidade": 5, "fonte": "PROPRIA", "codigo_composicao": "C1", "estado": "ce"})
    assert ok.status_code == 200, ok.text
    assert ok.json()["preco_unitario"] == 70.0
    assert ok.json()["preco_total"] == 350.0

    assert a.put(url, json={"codigo_composicao": "OUTRA"}).status_code == 400
    assert a.put(url, json={"fonte": "SINAPI"}).status_code == 400
    assert a.put(url, json={"estado": "SP"}).status_code == 400


def test_nao_converte_item_comum_em_proprio(a, banco, cenario):
    banco.tabelas["orcamento_itens"] = [
        {
            "id": "i1",
            "orcamento_id": cenario["orcamento"],
            "codigo_composicao": "X",
            "fonte": "SINAPI",
            "estado": "ce",
            "quantidade": 1,
            "preco_unitario": 5.0,
        }
    ]
    r = a.put(f"/orcamentos/{cenario['orcamento']}/itens/i1", json={"fonte": "PROPRIA"})
    assert r.status_code == 400


# --- Busca para o autocomplete -----------------------------------------------


def test_busca_composicoes_proprias_com_preco_do_orcamento(a, cenario):
    nova_composicao(a, cenario["base"], "OUTRA", [])
    url = f"/orcamentos/{cenario['orcamento']}/buscar-composicoes-proprias"

    achadas = a.get(f"{url}?termo=C1").json()
    assert len(achadas) == 1
    assert achadas[0]["codigo"] == "C1"
    assert achadas[0]["base_propria_id"] == cenario["base"]
    assert achadas[0]["base_propria_nome"] == "Minha base"
    assert achadas[0]["unidade"] == "M2"
    assert achadas[0]["preco_unitario"] == 70.0

    assert len(a.get(url).json()) == 2
    assert len(a.get(f"{url}?limite=1").json()) == 1
    assert a.get(f"{url}?limite=0").status_code == 422


# --- Exportações -------------------------------------------------------------


def test_excel_mostra_fonte_e_nome_da_base_propria(a, cenario):
    incluir(a, cenario)
    r = a.get(f"/orcamentos/{cenario['orcamento']}/excel")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws["H4"].value == "Fonte"
    assert ws["H5"].value == "Própria: Minha base"


def test_excel_de_item_comum_mostra_a_fonte():
    from app.modules.orcamento.export import gerar_planilha_orcamento_excel

    conteudo = gerar_planilha_orcamento_excel(
        {"nome": "x"}, [{"quantidade": 1, "preco_unitario": 2, "fonte": "SINAPI"}]
    )
    assert openpyxl.load_workbook(io.BytesIO(conteudo)).active["H5"].value == "SINAPI"


def test_pdf_do_orcamento_com_item_proprio(a, cenario):
    incluir(a, cenario)
    r = a.get(f"/orcamentos/{cenario['orcamento']}/pdf")
    assert r.status_code == 200
    assert r.content.startswith(b"%PDF")


def test_pdf_descreve_a_fonte_dos_itens_proprios():
    assert PdfService._descricao_item({"descricao": "Reboco", "fonte": "SINAPI"}) == "Reboco"
    assert (
        PdfService._descricao_item({"descricao": "Reboco", "fonte": "PROPRIA", "base_propria_nome": "Minha base"})
        == "Reboco (Fonte: Minha base)"
    )
    assert PdfService._descricao_item({"descricao": "Reboco", "fonte": "PROPRIA"}) == "Reboco (Fonte: Base própria)"

    itens = [
        {"descricao": "Reboco", "quantidade": 1, "preco_unitario": 10, "fonte": "PROPRIA", "base_propria_nome": "B"}
    ]
    orcamento = {"nome": "O", "cliente": "C", "created_at": "2026-01-01", "bdi": 0}
    assert PdfService().gerar_pdf(orcamento, itens)[:4] == b"%PDF"
    etapas = [{"id": "e1", "nome": "Etapa", "parent_id": None}]
    assert PdfService().gerar_pdf(orcamento, [{**itens[0], "etapa_id": "e1"}, itens[0]], etapas)[:4] == b"%PDF"


# --- anexar_nome_base_propria e regressão SINAPI -----------------------------


def test_anexar_nome_ignora_itens_de_outras_fontes_e_base_removida():
    servico = MagicMock()
    servico.repo.buscar_base.return_value = None
    itens = [{"fonte": "SINAPI"}, {"fonte": "PROPRIA", "base_propria_id": "b1"}, {"fonte": "PROPRIA"}]
    anexar_nome_base_propria(itens, servico, "u")
    assert "base_propria_nome" not in itens[0]
    assert itens[1]["base_propria_nome"] == "Base própria"
    assert itens[2]["base_propria_nome"] == "Base própria"
    assert anexar_nome_base_propria([{"fonte": "PROPRIA", "base_propria_id": "b"}], None, "u")[0]["base_propria_nome"]
    assert anexar_nome_base_propria([], servico, "u") == []


def test_fluxo_sinapi_nao_envia_base_propria_id(
    orcamento_item_service, orcamento_repo_mock, item_repo_mock, orcamento_item_repo_mock
):
    orcamento_repo_mock.buscar_por_id.return_value = {
        "id": "orc1",
        "estado": "SP",
        "base_referencia": MES,
        "tipo_composicao": TIPO,
        "bdi": 0.0,
    }
    item_repo_mock.buscar_por_codigo.return_value = [{"descricao": "Tijolo", "unidade": "un"}]
    item_repo_mock.buscar_preco.return_value = 10.0
    orcamento_item_repo_mock.criar.return_value = {"id": "i1", "quantidade": 2.0}
    orcamento_item_repo_mock.calcular_total_itens.return_value = 20.0

    orcamento_item_service.adicionar_item(
        "orc1", OrcamentoItemCreate(codigo_composicao="1", descricao="Tijolo", quantidade=2, unidade="un")
    )

    dados = orcamento_item_repo_mock.criar.call_args[0][0]
    assert "base_propria_id" not in dados
    assert dados["fonte"] == "SINAPI"


def test_servico_sem_base_propria_recusa_item_proprio(orcamento_item_service, orcamento_repo_mock):
    orcamento_repo_mock.buscar_por_id.return_value = {"id": "o", "estado": "CE", "user_id": "u"}
    with pytest.raises(ValueError, match="não disponível"):
        orcamento_item_service.adicionar_item(
            "o",
            OrcamentoItemCreate(
                codigo_composicao="C", descricao="d", quantidade=1, unidade="u", fonte="PROPRIA", base_propria_id="b"
            ),
        )
    with pytest.raises(ValueError, match="não disponível"):
        orcamento_item_service.atualizar_precos_propria("o")
    with pytest.raises(ValueError, match="não disponível"):
        orcamento_item_service.buscar_composicoes_proprias("o", "x")


def test_orcamento_inexistente_nas_novas_operacoes():
    repo = MagicMock()
    repo.buscar_por_id.return_value = None
    servico = OrcamentoItemService(MagicMock(), repo, MagicMock(), None, MagicMock())
    for chamada in (
        lambda: servico.atualizar_precos_propria("x"),
        lambda: servico.buscar_composicoes_proprias("x", ""),
    ):
        with pytest.raises(ValueError, match="Orçamento não encontrado"):
            chamada()
